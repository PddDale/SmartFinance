import json
import os
import tempfile
import unittest
from datetime import date, timedelta
from unittest.mock import patch

from Recursos.Codigo.app import app
from Recursos.Codigo import database, models


class TestSmartFinance(unittest.TestCase):
    def setUp(self):
        self.app = app.test_client()
        self.app.testing = True

    def test_01_index_dashboard(self):
        response = self.app.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SmartFinance", response.data)
        self.assertIn(b"Resumo Financeiro", response.data)
        self.assertIn(b"Patrim", response.data)
        self.assertIn(b"Limite e Uso por Categoria", response.data)
        self.assertIn(b"Porcentagem de uso:", response.data)
        self.assertIn(b"Gasto atual:", response.data)
        self.assertIn(b"Meta:", response.data)
        self.assertIn(b"Acima do limite:", response.data)
        self.assertIn(b"R$ ", response.data)

    def test_02_lancamentos_page(self):
        response = self.app.get("/lancamentos")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SmartFinance", response.data)
        self.assertIn(b"Gastos", response.data)

    def test_03_parcelas_page(self):
        response = self.app.get("/parcelas")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SmartFinance", response.data)
        self.assertIn(b"Compras Parceladas", response.data)

    def test_04_recorrentes_page(self):
        response = self.app.get("/recorrentes")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SmartFinance", response.data)
        self.assertIn(b"Gastos Recorrentes", response.data)

    def test_05_cartoes_page(self):
        response = self.app.get("/cartoes")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SmartFinance", response.data)
        self.assertIn(b"Cart", response.data)

    def test_06_anual_page(self):
        response = self.app.get("/anual")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SmartFinance", response.data)
        self.assertIn(b"Vis", response.data)
        conn = database.get_db_connection()
        categoria = conn.execute(
            "SELECT categoria, cor FROM limites_categoria ORDER BY categoria LIMIT 1;"
        ).fetchone()
        conn.close()
        self.assertIn(
            f"background-color: {categoria['cor']};".encode(),
            response.data,
        )

    def test_07_configuracoes_page(self):
        response = self.app.get("/configuracoes")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"SmartFinance", response.data)
        self.assertIn(b"Configura", response.data)
        conn = database.get_db_connection()
        categoria = conn.execute(
            "SELECT categoria, cor FROM limites_categoria ORDER BY categoria LIMIT 1;"
        ).fetchone()
        conn.close()
        self.assertIn(
            f"background-color: {categoria['cor']};".encode(),
            response.data,
        )

    def test_08_api_calcular_fatura(self):
        conn = database.get_db_connection()
        cartao = conn.execute("SELECT id FROM cartoes LIMIT 1;").fetchone()
        conn.close()
        self.assertIsNotNone(cartao)

        response = self.app.post(
            "/api/calcular-fatura",
            data=json.dumps({"data": "2026-10-26", "cartao_id": cartao["id"]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        data = json.loads(response.data)
        self.assertTrue(data["success"])
        self.assertIn("fatura_mes", data)
        self.assertIn("data_vencimento", data)

    def test_09_exportar_json(self):
        response = self.app.get("/configuracoes/exportar-json")
        self.assertEqual(response.status_code, 200)
        data = json.loads(response.data)
        self.assertIn("configuracoes", data)
        self.assertIn("rendas_anuais", data["configuracoes"])
        self.assertIn("saldos", data)
        self.assertIn("cartoes", data)
        self.assertIn("lancamentos", data)

    def test_10_novo_lancamento_and_delete(self):
        response = self.app.post(
            "/lancamentos/novo",
            data={
                "data": "2026-10-05",
                "descricao": "Teste Automatizado Lanche",
                "categoria": "Alimentação",
                "valor": "35,50",
                "metodo_pagamento": "Pix",
                "observacao": "Teste unitário",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Teste Automatizado Lanche", response.data)

        conn = database.get_db_connection()
        row = conn.execute(
            "SELECT id FROM lancamentos WHERE descricao = 'Teste Automatizado Lanche';"
        ).fetchone()
        conn.close()
        self.assertIsNotNone(row)

        response = self.app.post(
            f"/lancamentos/excluir/{row['id']}", follow_redirects=True
        )
        self.assertEqual(response.status_code, 200)

    @patch("os._exit")
    def test_11_api_encerrar(self, mock_exit):
        response = self.app.post("/api/encerrar")
        self.assertEqual(response.status_code, 200)
        data = json.loads(response.data)
        self.assertTrue(data["success"])
        self.assertIn("SmartFinance", data["message"])

    def test_12_api_categorizar_regras(self):
        casos = [
            ("iFood almoço executivo", "Alimentação"),
            ("Uber corrida aeroporto", "Transporte"),
            ("Netflix assinatura mensal", "Lazer"),
            ("Drogasil remédio pressão", "Saúde"),
            ("Conta de luz CEMIG", "Moradia"),
        ]
        for descricao, categoria_esperada in casos:
            with self.subTest(descricao=descricao):
                response = self.app.post(
                    "/api/categorizar",
                    data=json.dumps({"descricao": descricao}),
                    content_type="application/json",
                )
                self.assertEqual(response.status_code, 200)
                data = json.loads(response.data)
                self.assertTrue(data["success"])
                self.assertEqual(data["categoria"], categoria_esperada)

    def test_13_api_categorizar_historico(self):
        self.app.post(
            "/lancamentos/novo",
            data={
                "data": "2026-10-01",
                "descricao": "Mercadinho do Zé",
                "categoria": "Alimentação",
                "valor": "50.00",
                "metodo_pagamento": "Pix",
            },
        )
        response = self.app.post(
            "/api/categorizar",
            data=json.dumps({"descricao": "Mercadinho do Zé"}),
            content_type="application/json",
        )
        data = json.loads(response.data)
        self.assertTrue(data["success"])
        self.assertEqual(data["categoria"], "Alimentação")

        conn = database.get_db_connection()
        conn.execute("DELETE FROM lancamentos WHERE descricao='Mercadinho do Zé';")
        conn.commit()
        conn.close()

    def test_14_renda_anual_e_preferencias_persistem_por_ano(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                config_inicial = conn.execute(
                    "SELECT ano_ativo FROM configuracoes WHERE id = 1;"
                ).fetchone()
                conn.close()
                self.assertEqual(config_inicial["ano_ativo"], date.today().year)

                response = self.app.post(
                    "/configuracoes/salvar",
                    data={
                        "ano_ativo": "2032",
                        "modo_renda": "anual",
                        "salario_anual": "120000",
                        "outras_rendas_anuais": "24000",
                        "saldo_conta_corrente": "2500",
                        "valor_investido": "15000",
                        "formato_data": "semana_extenso",
                    },
                )
                self.assertEqual(response.status_code, 302)

                conn = database.get_db_connection()
                config = conn.execute(
                    "SELECT ano_ativo, formato_data FROM configuracoes WHERE id = 1;"
                ).fetchone()
                renda = conn.execute(
                    "SELECT salario_mensal, outras_rendas_mensais "
                    "FROM rendas_anuais WHERE ano = 2032;"
                ).fetchone()
                conn.close()
                self.assertEqual(config["ano_ativo"], 2032)
                self.assertEqual(config["formato_data"], "semana_extenso")
                self.assertEqual(renda["salario_mensal"], 10000)
                self.assertEqual(renda["outras_rendas_mensais"], 2000)

                response = self.app.get("/")
                dias = [
                    "segunda-feira",
                    "terça-feira",
                    "quarta-feira",
                    "quinta-feira",
                    "sexta-feira",
                    "sábado",
                    "domingo",
                ]
                trecho_data = (
                    f"{dias[date.today().weekday()]}, {date.today().day} de "
                )
                self.assertIn(trecho_data.encode("utf-8"), response.data)

                self.app.post("/configuracoes/ano", data={"ano_ativo": "2033"})
                conn = database.get_db_connection()
                renda_ano_novo = conn.execute(
                    "SELECT ano FROM rendas_anuais WHERE ano = 2033;"
                ).fetchone()
                conn.close()
                self.assertIsNotNone(renda_ano_novo)

    def test_15_renda_variavel_mensal_e_limite_zero_no_dashboard(self):
        self.assertEqual(models.formatar_moeda(1000), "R$ 1.000,00")
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                self.assertFalse(models.get_configuracoes()["renda_variavel_mensal"])
                dados = {
                    "ano_ativo": "2034",
                    "modo_renda": "mensal",
                    "salario_mensal": "100",
                    "outras_rendas_mensais": "0",
                    "saldo_conta_corrente": "0",
                    "valor_investido": "0",
                    "formato_data": "dd/mm/aaaa",
                    "renda_variavel_mensal": "1",
                    "salario_mes_01": "1000",
                    "outras_rendas_mes_01": "0",
                    "salario_mes_02": "2500",
                    "outras_rendas_mes_02": "0",
                }
                response = self.app.post("/configuracoes/salvar", data=dados)
                self.assertEqual(response.status_code, 302)
                self.assertEqual(models.get_resumo_mensal("2034-01")["renda_total"], 1000)
                self.assertEqual(models.get_resumo_mensal("2034-02")["renda_total"], 2500)
                self.assertEqual(
                    models.get_visao_anual(2034)["total_anual_renda"], 4500
                )
                models.update_limite_categoria("Moradia", 0)
                categorias_painel = models.get_resumo_mensal("2034-01")[
                    "categorias_dashboard"
                ]
                self.assertNotIn("Moradia", [cat["categoria"] for cat in categorias_painel])

    def test_16_importacao_excel_mapeia_colunas_e_valida_moeda_brasileira(self):
        from io import BytesIO
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["Compra", "Preço", "Data"])
                sheet.append(["Mercado", "R$ 1.234,50", "05/10/2026"])
                arquivo = BytesIO()
                workbook.save(arquivo)
                arquivo_bytes = arquivo.getvalue()
                response = self.app.post(
                    "/lancamentos/importar/previa",
                    data={"arquivo": (BytesIO(arquivo_bytes), "compras.xlsx")},
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(json.loads(response.data)["total_linhas"], 1)

                response = self.app.post(
                    "/lancamentos/importar",
                    data={
                        "arquivo": (BytesIO(arquivo_bytes), "compras.xlsx"),
                        "data_col": "2",
                        "descricao_col": "0",
                        "valor_col": "1",
                        "categoria": "Alimentação",
                        "metodo_pagamento": "Pix",
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(json.loads(response.data)["quantidade"], 1)
                conn = database.get_db_connection()
                lancamento = conn.execute(
                    "SELECT data, valor FROM lancamentos WHERE descricao = 'Mercado';"
                ).fetchone()
                conn.close()
                self.assertEqual(lancamento["data"], "2026-10-05")
                self.assertEqual(lancamento["valor"], 1234.5)

    def test_16a_importacao_excel_mapeia_categoria_por_linha(self):
        from io import BytesIO
        from openpyxl import Workbook

        def planilha_excel(linhas):
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Compra", "Preço", "Data", "Categoria"])
            for linha in linhas:
                sheet.append(linha)
            arquivo = BytesIO()
            workbook.save(arquivo)
            return arquivo.getvalue()

        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                arquivo_bytes = planilha_excel([
                    ["Farmácia", "25,50", "05/10/2026", "Saúde"],
                    ["Mercado", "80,00", "06/10/2026", "Alimentação"],
                ])
                response = self.app.post(
                    "/lancamentos/importar",
                    data={
                        "arquivo": (BytesIO(arquivo_bytes), "compras.xlsx"),
                        "data_col": "2",
                        "descricao_col": "0",
                        "valor_col": "1",
                        "categoria_col": "3",
                        "categoria": "Outros",
                        "metodo_pagamento": "Pix",
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(json.loads(response.data)["quantidade"], 2)

                conn = database.get_db_connection()
                categorias_importadas = dict(conn.execute(
                    "SELECT descricao, categoria FROM lancamentos;"
                ).fetchall())
                conn.close()
                self.assertEqual(categorias_importadas, {
                    "Farmácia": "Saúde",
                    "Mercado": "Alimentação",
                })

                arquivo_invalido = planilha_excel([
                    ["Compra válida", "10,00", "07/10/2026", "Moradia"],
                    ["Categoria incorreta", "15,00", "07/10/2026", "Não cadastrada"],
                ])
                response = self.app.post(
                    "/lancamentos/importar",
                    data={
                        "arquivo": (BytesIO(arquivo_invalido), "compras.xlsx"),
                        "data_col": "2",
                        "descricao_col": "0",
                        "valor_col": "1",
                        "categoria_col": "3",
                        "categoria": "Outros",
                        "metodo_pagamento": "Pix",
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 400)
                self.assertIn(b"Linha 3", response.data)
                conn = database.get_db_connection()
                total = conn.execute(
                    "SELECT COUNT(*) FROM lancamentos;"
                ).fetchone()[0]
                conn.close()
                self.assertEqual(total, 2)

    def test_16b_importacao_excel_aceita_padrao_decimal_selecionado(self):
        from io import BytesIO
        from openpyxl import Workbook

        def arquivo_com_valor(valor):
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Compra", "Preço", "Data"])
            sheet.append(["Importação decimal", valor, "05/10/2026"])
            arquivo = BytesIO()
            workbook.save(arquivo)
            return arquivo.getvalue()

        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                for valor_planilha, formato, valor_esperado in (
                    ("1.234,56", "brasileiro", 1234.56),
                    ("1,234.56", "internacional", 1234.56),
                ):
                    response = self.app.post(
                        "/lancamentos/importar",
                        data={
                            "arquivo": (
                                BytesIO(arquivo_com_valor(valor_planilha)),
                                "compras.xlsx",
                            ),
                            "data_col": "2",
                            "descricao_col": "0",
                            "valor_col": "1",
                            "categoria": "Outros",
                            "metodo_pagamento": "Pix",
                            "formato_decimal": formato,
                        },
                        content_type="multipart/form-data",
                    )
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(
                        json.loads(response.data)["quantidade"], 1
                    )
                    conn = database.get_db_connection()
                    importado = conn.execute(
                        "SELECT valor FROM lancamentos "
                        "WHERE descricao = 'Importação decimal';"
                    ).fetchone()
                    conn.close()
                    self.assertAlmostEqual(importado["valor"], valor_esperado)

                response = self.app.post(
                    "/lancamentos/importar",
                    data={
                        "arquivo": (
                            BytesIO(arquivo_com_valor("1,234.56")),
                            "compras.xlsx",
                        ),
                        "data_col": "2",
                        "descricao_col": "0",
                        "valor_col": "1",
                        "categoria": "Outros",
                        "metodo_pagamento": "Pix",
                        "formato_decimal": "brasileiro",
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 400)
                self.assertIn(b"brasileiro", response.data)

    def test_16c_importacao_excel_aceita_categoria_com_emoji(self):
        from io import BytesIO
        from datetime import datetime
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Data", "Descrição", "Categoria  ▾", "Valor (R$)"])
        sheet.append([datetime(2026, 5, 19), "Restaurante", "🍔 Alimentação", 7.18])
        arquivo = BytesIO()
        workbook.save(arquivo)

        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                response = self.app.post(
                    "/lancamentos/importar",
                    data={
                        "arquivo": (BytesIO(arquivo.getvalue()), "Book1.xlsx"),
                        "data_col": "0",
                        "descricao_col": "1",
                        "categoria_col": "2",
                        "valor_col": "3",
                        "categoria": "Outros",
                        "metodo_pagamento": "Pix",
                        "formato_decimal": "internacional",
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200)
                conn = database.get_db_connection()
                lancamento = conn.execute(
                    "SELECT data, categoria, valor FROM lancamentos "
                    "WHERE descricao = 'Restaurante';"
                ).fetchone()
                conn.close()
                self.assertEqual(lancamento["data"], "2026-05-19")
                self.assertEqual(lancamento["categoria"], "Alimentação")
                self.assertAlmostEqual(lancamento["valor"], 7.18)

    def test_17_pagamentos_de_parcela_e_recorrencia_geram_um_lancamento(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                cartao_id = conn.execute("""
                    INSERT INTO cartoes (nome, limite_total, fechamento_dia, vencimento_dia)
                    VALUES ('Cartão teste', 1000, 25, 5);
                """).lastrowid
                conn.commit()
                conn.close()
                models.add_compra_parcelada(
                    "Notebook teste", date.today().isoformat(), 300, 0, 3,
                    "Tecnologia", "Cartão de Crédito", cartao_id,
                )
                conn = database.get_db_connection()
                parcela = conn.execute(
                    "SELECT id FROM parcelas_detalhe ORDER BY numero_parcela LIMIT 1;"
                ).fetchone()
                compra = conn.execute(
                    "SELECT cartao_id FROM compras_parceladas LIMIT 1;"
                ).fetchone()
                mes_parcela = conn.execute(
                    "SELECT ano_mes FROM parcelas_detalhe WHERE id = ?;",
                    (parcela["id"],),
                ).fetchone()["ano_mes"]
                self.assertEqual(compra["cartao_id"], cartao_id)
                conn.close()
                limite_antes = models.get_cartoes()[0]["limite_comprometido"]
                self.assertGreater(limite_antes, 0)

                resposta = self.app.post(
                    f"/parcelas/{parcela['id']}/pagar",
                    data={"metodo_pagamento": "Pix"},
                )
                self.assertEqual(resposta.status_code, 302)
                self.app.post(
                    f"/parcelas/{parcela['id']}/pagar",
                    data={"metodo_pagamento": "Pix"},
                )
                models.add_gasto_recorrente(
                    "Internet teste", 80, "Moradia", 5, "Cartão de Débito"
                )
                conn = database.get_db_connection()
                recorrente_id = conn.execute(
                    "SELECT id FROM gastos_recorrentes WHERE descricao = 'Internet teste';"
                ).fetchone()["id"]
                conn.close()
                ano_mes = date.today().strftime("%Y-%m")
                self.app.post(
                    f"/recorrentes/{recorrente_id}/pagar",
                    data={"ano_mes": ano_mes, "metodo_pagamento": "Dinheiro"},
                )
                conn = database.get_db_connection()
                total = conn.execute(
                    "SELECT COUNT(*) FROM lancamentos WHERE descricao IN "
                    "('Notebook teste', 'Internet teste');"
                ).fetchone()[0]
                parcela_pago = conn.execute(
                    "SELECT status, lancamento_id, metodo_pagamento_pago "
                    "FROM parcelas_detalhe WHERE id = ?;",
                    (parcela["id"],),
                ).fetchone()
                recorrente_pago = conn.execute(
                    "SELECT pago, lancamento_id FROM recorrentes_status_mes "
                    "WHERE recorrente_id = ? AND ano_mes = ?;",
                    (recorrente_id, ano_mes),
                ).fetchone()
                conn.close()
                self.assertEqual(total, 2)
                self.assertEqual(parcela_pago["status"], "Pago")
                self.assertIsNotNone(parcela_pago["lancamento_id"])
                self.assertEqual(parcela_pago["metodo_pagamento_pago"], "Pix")
                self.assertEqual(recorrente_pago["pago"], 1)
                self.assertIsNotNone(recorrente_pago["lancamento_id"])
                self.assertLess(
                    models.get_cartoes()[0]["limite_comprometido"], limite_antes
                )
                self.assertEqual(
                    models.get_resumo_mensal(mes_parcela)["total_gastos"], 180
                )
                self.assertEqual(
                    models.get_resumo_mensal(ano_mes)["total_gastos"], 80
                )

                conn = database.get_db_connection()
                lancamento_parcela_id = parcela_pago["lancamento_id"]
                lancamento_recorrente_id = recorrente_pago["lancamento_id"]
                conn.close()

                self.app.post(f"/lancamentos/excluir/{lancamento_parcela_id}")
                conn = database.get_db_connection()
                parcela_pendente = conn.execute(
                    "SELECT status, lancamento_id FROM parcelas_detalhe WHERE id = ?;",
                    (parcela["id"],),
                ).fetchone()
                self.assertEqual(parcela_pendente["status"], "Pendente")
                self.assertIsNone(parcela_pendente["lancamento_id"])
                self.assertIsNone(conn.execute(
                    "SELECT id FROM lancamentos WHERE id = ?;",
                    (lancamento_parcela_id,),
                ).fetchone())
                conn.close()

                self.app.post(
                    f"/parcelas/{parcela['id']}/pagar",
                    data={"metodo_pagamento": "Pix"},
                )
                conn = database.get_db_connection()
                lancamento_novo_id = conn.execute(
                    "SELECT lancamento_id FROM parcelas_detalhe WHERE id = ?;",
                    (parcela["id"],),
                ).fetchone()["lancamento_id"]
                conn.close()
                self.app.post(
                    f"/parcelas/{parcela['id']}/desfazer-pagamento"
                )
                conn = database.get_db_connection()
                parcela_desfeita = conn.execute(
                    "SELECT status, lancamento_id FROM parcelas_detalhe WHERE id = ?;",
                    (parcela["id"],),
                ).fetchone()
                self.assertEqual(parcela_desfeita["status"], "Pendente")
                self.assertIsNone(parcela_desfeita["lancamento_id"])
                self.assertIsNone(conn.execute(
                    "SELECT id FROM lancamentos WHERE id = ?;",
                    (lancamento_novo_id,),
                ).fetchone())
                conn.close()

                self.app.post(f"/lancamentos/excluir/{lancamento_recorrente_id}")
                conn = database.get_db_connection()
                recorrente_pendente = conn.execute(
                    "SELECT pago, lancamento_id FROM recorrentes_status_mes "
                    "WHERE recorrente_id = ? AND ano_mes = ?;",
                    (recorrente_id, ano_mes),
                ).fetchone()
                self.assertEqual(recorrente_pendente["pago"], 0)
                self.assertIsNone(recorrente_pendente["lancamento_id"])
                conn.close()

    def test_18_recursos_estaticos_e_banco_local(self):
        css = self.app.get("/static/css/style.css")
        js = self.app.get("/static/js/app.js")
        try:
            self.assertEqual(css.status_code, 200)
            self.assertEqual(js.status_code, 200)
        finally:
            css.close()
            js.close()
        self.assertTrue(
            database.DB_PATH.endswith(
                os.path.join("Recursos", "Dados", "finance.db")
            )
        )

    def test_19_migracao_reabre_pagamentos_com_lancamento_orfao(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_compra_parcelada(
                    "Compra legada", date.today().isoformat(), 90, 0, 1,
                    "Outros", "Pix", None,
                )
                conn = database.get_db_connection()
                parcela_id = conn.execute(
                    "SELECT id FROM parcelas_detalhe LIMIT 1;"
                ).fetchone()["id"]
                conn.close()
                models.registrar_pagamento_parcela(parcela_id, "Pix")
                models.add_gasto_recorrente(
                    "Recorrente legado", 20, "Outros", 1, "Pix"
                )
                conn = database.get_db_connection()
                recorrente_id = conn.execute(
                    "SELECT id FROM gastos_recorrentes WHERE descricao = 'Recorrente legado';"
                ).fetchone()["id"]
                conn.close()
                ano_mes = date.today().strftime("%Y-%m")
                models.registrar_pagamento_recorrente(
                    recorrente_id, ano_mes, "Pix"
                )
                conn = database.get_db_connection()
                lancamento_id = conn.execute(
                    "SELECT lancamento_id FROM parcelas_detalhe WHERE id = ?;",
                    (parcela_id,),
                ).fetchone()["lancamento_id"]
                conn.execute("DELETE FROM lancamentos WHERE id = ?;", (lancamento_id,))
                lancamento_recorrente_id = conn.execute(
                    "SELECT lancamento_id FROM recorrentes_status_mes "
                    "WHERE recorrente_id = ? AND ano_mes = ?;",
                    (recorrente_id, ano_mes),
                ).fetchone()["lancamento_id"]
                conn.execute(
                    "DELETE FROM lancamentos WHERE id = ?;",
                    (lancamento_recorrente_id,),
                )
                conn.commit()
                conn.close()

                database.init_db()
                conn = database.get_db_connection()
                parcela = conn.execute(
                    "SELECT status, lancamento_id FROM parcelas_detalhe WHERE id = ?;",
                    (parcela_id,),
                ).fetchone()
                recorrente = conn.execute(
                    "SELECT pago, lancamento_id FROM recorrentes_status_mes "
                    "WHERE recorrente_id = ? AND ano_mes = ?;",
                    (recorrente_id, ano_mes),
                ).fetchone()
                conn.close()
                self.assertEqual(parcela["status"], "Pendente")
                self.assertIsNone(parcela["lancamento_id"])
                self.assertEqual(recorrente["pago"], 0)
                self.assertIsNone(recorrente["lancamento_id"])

    def test_19_exclusao_por_periodo_preserva_cadastros_e_outras_parcelas(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                cartao_id = conn.execute("""
                    INSERT INTO cartoes
                        (nome, limite_total, fechamento_dia, vencimento_dia)
                    VALUES ('Cartão de teste', 1000, 10, 20);
                """).lastrowid
                compra_id = conn.execute("""
                    INSERT INTO compras_parceladas (
                        descricao, data_compra, valor_total, entrada, num_parcelas,
                        valor_parcela, categoria, metodo_pagamento, cartao_id,
                        mes_inicio, mes_fim
                    ) VALUES (
                        'Compra de teste', '2026-10-01', 300, 0, 3, 100,
                        'Outros', 'Cartão de Crédito', ?, '2026-10', '2026-12'
                    );
                """, (cartao_id,)).lastrowid
                conn.executemany("""
                    INSERT INTO parcelas_detalhe (
                        compra_id, numero_parcela, total_parcelas, ano_mes,
                        valor, data_vencimento, status
                    ) VALUES (?, ?, 3, ?, 100, ?, 'Pendente');
                """, [
                    (compra_id, 1, "2026-10", "2026-10-05"),
                    (compra_id, 2, "2026-11", "2026-11-05"),
                ])
                conn.executemany("""
                    INSERT INTO lancamentos (
                        data, ano_mes, descricao, categoria, valor, metodo_pagamento
                    ) VALUES (?, ?, ?, 'Outros', 10, 'Pix');
                """, [
                    ("2026-10-05", "2026-10", "Excluir"),
                    ("2026-10-06", "2026-10", "Preservar"),
                ])
                conn.execute("""
                    INSERT INTO gastos_recorrentes (descricao, valor_mensal, categoria)
                    VALUES ('Assinatura de teste', 20, 'Outros');
                """)
                conn.execute("""
                    UPDATE configuracoes
                    SET saldo_conta_corrente = 1234, valor_investido = 5678
                    WHERE id = 1;
                """)
                conn.commit()
                conn.close()

                response = self.app.post(
                    "/configuracoes/excluir-dados",
                    data={"tipo": "dia", "referencia": "data-invalida"},
                )
                self.assertEqual(response.status_code, 302)

                response = self.app.post(
                    "/configuracoes/excluir-dados",
                    data={"tipo": "dia", "referencia": "2026-10-05"},
                )
                self.assertEqual(response.status_code, 302)

                conn = database.get_db_connection()
                self.assertEqual(
                    [tuple(row) for row in conn.execute(
                        "SELECT descricao FROM lancamentos;"
                    ).fetchall()],
                    [("Preservar",)],
                )
                self.assertEqual(
                    conn.execute("SELECT COUNT(*) FROM parcelas_detalhe;").fetchone()[0],
                    1,
                )
                self.assertIsNotNone(conn.execute(
                    "SELECT id FROM compras_parceladas WHERE id = ?;", (compra_id,)
                ).fetchone())
                self.assertEqual(conn.execute("SELECT COUNT(*) FROM cartoes;").fetchone()[0], 1)
                self.assertEqual(
                    conn.execute("SELECT COUNT(*) FROM gastos_recorrentes;").fetchone()[0],
                    1,
                )
                config = conn.execute(
                    "SELECT saldo_conta_corrente, valor_investido "
                    "FROM configuracoes WHERE id = 1;"
                ).fetchone()
                conn.close()
                self.assertEqual(tuple(config), (1234, 5678))

    def test_20_exclusao_total_apaga_dados_e_zera_saldos(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                conn.execute("""
                    INSERT INTO cartoes
                        (nome, limite_total, fechamento_dia, vencimento_dia)
                    VALUES ('Cartão de teste', 1000, 10, 20);
                """)
                conn.execute("""
                    INSERT INTO lancamentos (
                        data, ano_mes, descricao, categoria, valor, metodo_pagamento
                    ) VALUES ('2026-10-05', '2026-10', 'Gasto de teste', 'Outros', 10, 'Pix');
                """)
                conn.execute("""
                    INSERT INTO gastos_recorrentes (descricao, valor_mensal, categoria)
                    VALUES ('Assinatura de teste', 20, 'Outros');
                """)
                conn.execute("""
                    UPDATE configuracoes
                    SET salario = 5000, outras_rendas = 500,
                        saldo_conta_corrente = 1234, valor_investido = 5678
                    WHERE id = 1;
                """)
                conn.commit()
                conn.close()

                resposta_sem_confirmacao = self.app.post(
                    "/configuracoes/excluir-dados",
                    data={"tipo": "todos"},
                )
                self.assertEqual(resposta_sem_confirmacao.status_code, 302)
                conn = database.get_db_connection()
                self.assertEqual(conn.execute("SELECT COUNT(*) FROM lancamentos;").fetchone()[0], 1)
                conn.close()

                response = self.app.post(
                    "/configuracoes/excluir-dados",
                    data={"tipo": "todos", "confirmacao": "EXCLUIR TUDO"},
                )
                self.assertEqual(response.status_code, 302)

                conn = database.get_db_connection()
                for tabela in (
                    "lancamentos",
                    "parcelas_detalhe",
                    "compras_parceladas",
                    "recorrentes_status_mes",
                    "gastos_recorrentes",
                    "cartoes",
                ):
                    with self.subTest(tabela=tabela):
                        self.assertEqual(
                            conn.execute(f"SELECT COUNT(*) FROM {tabela};").fetchone()[0],
                            0,
                        )
                config = conn.execute("""
                    SELECT salario, outras_rendas, saldo_conta_corrente, valor_investido
                    FROM configuracoes WHERE id = 1;
                """).fetchone()
                self.assertEqual(tuple(config), (0, 0, 0, 0))
                self.assertEqual(
                    conn.execute("SELECT COUNT(*) FROM rendas_anuais;").fetchone()[0],
                    1,
                )
                self.assertEqual(
                    conn.execute(
                        "SELECT COUNT(*) FROM limites_categoria WHERE limite_mensal != 0;"
                    ).fetchone()[0],
                    0,
                )
                conn.close()

    def test_21_exclusao_por_semana_mes_e_ano_respeita_limites(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                conn.executemany("""
                    INSERT INTO lancamentos (
                        data, ano_mes, descricao, categoria, valor, metodo_pagamento
                    ) VALUES (?, ?, ?, 'Outros', 10, 'Pix');
                """, [
                    ("2026-10-05", "2026-10", "Segunda"),
                    ("2026-10-11", "2026-10", "Domingo"),
                    ("2026-10-12", "2026-10", "Semana seguinte"),
                    ("2026-11-01", "2026-11", "Mês seguinte"),
                    ("2027-01-01", "2027-01", "Ano seguinte"),
                ])
                conn.commit()
                conn.close()

                response = self.app.post(
                    "/configuracoes/excluir-dados",
                    data={"tipo": "semana", "referencia": "2026-10-07"},
                )
                self.assertEqual(response.status_code, 302)
                conn = database.get_db_connection()
                restantes = {
                    row["descricao"] for row in conn.execute(
                        "SELECT descricao FROM lancamentos;"
                    ).fetchall()
                }
                conn.close()
                self.assertEqual(
                    restantes,
                    {"Semana seguinte", "Mês seguinte", "Ano seguinte"},
                )

                response = self.app.post(
                    "/configuracoes/excluir-dados",
                    data={"tipo": "mes", "referencia": "2026-11"},
                )
                self.assertEqual(response.status_code, 302)
                conn = database.get_db_connection()
                restantes = {
                    row["descricao"] for row in conn.execute(
                        "SELECT descricao FROM lancamentos;"
                    ).fetchall()
                }
                conn.close()
                self.assertEqual(restantes, {"Semana seguinte", "Ano seguinte"})

                response = self.app.post(
                    "/configuracoes/excluir-dados",
                    data={"tipo": "ano", "referencia": "2026"},
                )
                self.assertEqual(response.status_code, 302)
                conn = database.get_db_connection()
                restantes = {
                    row["descricao"] for row in conn.execute(
                        "SELECT descricao FROM lancamentos;"
                    ).fetchall()
                }
                conn.close()
                self.assertEqual(restantes, {"Ano seguinte"})

    def test_22_cores_de_categoria_consistentes_em_todas_as_telas(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                hoje = date.today()
                hoje_iso = hoje.isoformat()
                ano_mes = hoje.strftime("%Y-%m")

                conn = database.get_db_connection()
                categoria = conn.execute("""
                    SELECT categoria, cor FROM limites_categoria
                    WHERE categoria = 'Alimentação';
                """).fetchone()
                cartao_id = conn.execute("""
                    INSERT INTO cartoes
                        (nome, limite_total, fechamento_dia, vencimento_dia)
                    VALUES ('Cartão de teste', 1000, 10, 20);
                """).lastrowid
                compra_id = conn.execute("""
                    INSERT INTO compras_parceladas (
                        descricao, data_compra, valor_total, entrada, num_parcelas,
                        valor_parcela, categoria, metodo_pagamento, cartao_id,
                        mes_inicio, mes_fim
                    ) VALUES (
                        'Parcela de teste', ?, 100, 0, 1, 100,
                        'Alimentação', 'Cartão de Crédito', ?, ?, ?
                    );
                """, (hoje_iso, cartao_id, ano_mes, ano_mes)).lastrowid
                conn.execute("""
                    INSERT INTO parcelas_detalhe (
                        compra_id, numero_parcela, total_parcelas, ano_mes,
                        valor, data_vencimento, status
                    ) VALUES (?, 1, 1, ?, 100, ?, 'Pendente');
                """, (compra_id, ano_mes, hoje_iso))
                conn.execute("""
                    INSERT INTO lancamentos (
                        data, ano_mes, descricao, categoria, valor, metodo_pagamento,
                        cartao_id, mes_vencimento, data_vencimento
                    ) VALUES (?, ?, 'Compra teste', 'Alimentação', 50, 'Cartão de Crédito',
                              ?, ?, ?);
                """, (hoje_iso, ano_mes, cartao_id, ano_mes, hoje_iso))
                conn.execute("""
                    INSERT INTO gastos_recorrentes (
                        descricao, valor_mensal, categoria, cartao_id
                    ) VALUES ('Recorrência teste', 25, 'Alimentação', ?);
                """, (cartao_id,))
                conn.commit()
                conn.close()

                paginas = (
                    "/",
                    "/lancamentos",
                    "/recorrentes",
                    "/parcelas",
                    f"/cartoes?mes={ano_mes}",
                    f"/anual?ano={hoje.year}",
                    "/configuracoes",
                )
                for pagina in paginas:
                    with self.subTest(pagina=pagina):
                        response = self.app.get(pagina)
                        self.assertEqual(response.status_code, 200)
                        cor_esperada = (
                            categoria["cor"]
                            if pagina.startswith(("/anual", "/configuracoes"))
                            else f"{categoria['cor']}18"
                        )
                        self.assertIn(
                            f"background-color: {cor_esperada};".encode(),
                            response.data,
                        )

    def test_23_formatacao_de_mes_da_fatura_respeita_preferencia(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                conn.execute(
                    "UPDATE configuracoes SET formato_data = 'extenso' WHERE id = 1;"
                )
                conn.commit()
                conn.close()

                response = self.app.get("/cartoes?mes=2026-10")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"outubro de 2026", response.data)
                self.assertIn(b'id="fatura_mes" value="2026-10"', response.data)
                self.assertNotIn(b'type="month"', response.data)

    def test_24_valores_redefinidos_sao_exibidos_com_virgula(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.excluir_todos_dados()

                response = self.app.get("/configuracoes")
                self.assertEqual(response.status_code, 200)
                self.assertIn(
                    b'name="salario_mensal" data-income-key="salario" '
                    b'value="0,00"',
                    response.data,
                )
                self.assertIn(
                    b'name="saldo_conta_corrente" value="0,00"',
                    response.data,
                )

    def test_25_sugestoes_de_lancamentos_usam_descricoes_salvas(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_lancamento(
                    "2026-09-10", "Supermercado", "Alimentação", 42, "", "Pix"
                )
                models.add_lancamento(
                    "2026-09-11", "Supercola", "Outros", 5, "", "Dinheiro"
                )

                response = self.app.get("/api/lancamentos/sugestoes?q=Supe")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    json.loads(response.data)["sugestoes"],
                    ["Supercola", "Supermercado"],
                )

    def test_26_configuracoes_aceitam_valor_decimal_com_virgula(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                response = self.app.post(
                    "/configuracoes/salvar",
                    data={
                        "ano_ativo": str(date.today().year),
                        "modo_renda": "mensal",
                        "salario_mensal": "1.234,56",
                        "outras_rendas_mensais": "345,67",
                        "saldo_conta_corrente": "2.345,67",
                        "valor_investido": "0,00",
                        "formato_data": "dd/mm/aaaa",
                    },
                )
                self.assertEqual(response.status_code, 302)

                config = models.get_configuracoes()
                self.assertEqual(config["salario"], 1234.56)
                self.assertEqual(config["outras_rendas"], 345.67)
                self.assertEqual(config["saldo_conta_corrente"], 2345.67)

    def test_27_pagamento_recorrente_pendente_so_vira_lancamento_apos_confirmacao(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_gasto_recorrente(
                    "Internet teste", 80, "Moradia", date.today().day, "Pix"
                )
                conn = database.get_db_connection()
                recorrente_id = conn.execute(
                    "SELECT id FROM gastos_recorrentes WHERE descricao = ?;",
                    ("Internet teste",),
                ).fetchone()["id"]
                conn.close()

                pendentes = models.get_pagamentos_recorrentes_pendentes()
                self.assertTrue(any(
                    p["id"] == recorrente_id for p in pendentes
                ))
                proximo_mes = (date.today().replace(day=1) + timedelta(days=32))
                models.add_gasto_recorrente(
                    "Nova conta teste", 25, "Moradia", 1, "Pix",
                    mes_inicio=proximo_mes.strftime("%Y-%m"),
                )
                self.assertFalse(any(
                    p["descricao"] == "Nova conta teste"
                    for p in models.get_pagamentos_recorrentes_pendentes()
                ))
                conn = database.get_db_connection()
                self.assertEqual(
                    conn.execute(
                        "SELECT COUNT(*) FROM lancamentos WHERE descricao = ?;",
                        ("Internet teste",),
                    ).fetchone()[0],
                    0,
                )
                conn.close()

                response = self.app.get("/lancamentos")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"Pagamentos recorrentes a confirmar", response.data)
                self.assertIn(b"Internet teste", response.data)
                self.assertIn(b"data-recorrente-pendente=", response.data)
                self.assertIn(b"1 m\xc3\xaas(es) a confirmar", response.data)

                models.add_gasto_recorrente(
                    "Aluguel teste", 1200, "Moradia", 1, "Pix",
                    mes_inicio="2026-06",
                )
                conn = database.get_db_connection()
                recorrencia_retroativa = conn.execute(
                    "SELECT id FROM gastos_recorrentes WHERE descricao = ?;",
                    ("Aluguel teste",),
                ).fetchone()["id"]
                conn.close()
                response = self.app.get("/lancamentos")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"Aluguel teste", response.data)
                self.assertIn(b"5 m\xc3\xaas(es) a confirmar", response.data)
                self.assertIn(b"data-payment-month=\"2026-06\"", response.data)
                self.assertIn(b"data-payment-month=\"2026-10\"", response.data)
                self.assertIn(b"2 compra(s) pendente(s)", response.data)
                self.assertEqual(
                    response.data.count(b"data-recorrente-pendente="),
                    2,
                )
                self.assertEqual(
                    response.data.count(
                        f'data-recorrente-pendente="{recorrencia_retroativa}"'.encode()
                    ),
                    1,
                )

                response = self.app.post(
                    f"/recorrentes/{recorrente_id}/pagar",
                    data={
                        "ano_mes": date.today().strftime("%Y-%m"),
                        "metodo_pagamento": "Pix",
                    },
                )
                self.assertEqual(response.status_code, 302)
                self.assertEqual(response.location, "/lancamentos")

                conn = database.get_db_connection()
                self.assertEqual(
                    conn.execute(
                        "SELECT COUNT(*) FROM lancamentos WHERE descricao = ?;",
                        ("Internet teste",),
                    ).fetchone()[0],
                    1,
                )
                conn.close()
                response = self.app.get("/lancamentos")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"Internet teste", response.data)
                self.assertFalse(any(
                    p["id"] == recorrente_id
                    for p in models.get_pagamentos_recorrentes_pendentes()
                ))

    def test_28_recorrentes_exibem_status_mensal_do_ano(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_gasto_recorrente(
                    "Aluguel teste", 1200, "Moradia", 5, "Pix"
                )
                conn = database.get_db_connection()
                recorrente_id = conn.execute(
                    "SELECT id FROM gastos_recorrentes WHERE descricao = ?;",
                    ("Aluguel teste",),
                ).fetchone()["id"]
                conn.close()
                ano = date.today().year
                mes_atual = date.today().month
                ano_mes_pausado = f"{ano:04d}-{mes_atual:02d}"
                models.toggle_recorrente_mes(recorrente_id, ano_mes_pausado, False)
                conn = database.get_db_connection()
                conn.execute(
                    "UPDATE configuracoes SET formato_data = 'extenso' WHERE id = 1;"
                )
                conn.commit()
                conn.close()

                response = self.app.get(f"/recorrentes?ano={ano}")
                self.assertEqual(response.status_code, 200)
                self.assertIn(
                    f"Total Recorrente Ativo em {ano}".encode(), response.data
                )
                self.assertIn(
                    f'data-mes="{ano_mes_pausado}"'.encode(), response.data
                )
                meses_nomes = (
                    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
                    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
                )
                self.assertIn(
                    f'aria-label="Aluguel teste: {meses_nomes[mes_atual - 1]} de {ano}"'.encode(),
                    response.data,
                )
                self.assertIn(b"Pausado", response.data)
                self.assertIn(
                    f"janeiro de {ano}".encode(), response.data
                )
                self.assertIn(
                    f"junho de {ano}".encode(), response.data
                )
                self.assertNotIn(b'type="month"', response.data)
                self.assertIn(b'data-formatted-month-selector', response.data)

    def test_29_recorrente_com_inicio_passado_gera_cobrancas_pendentes_desde_entao(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_gasto_recorrente(
                    "Aluguel retroativo", 1500, "Moradia", 1, "Pix",
                    mes_inicio="2026-06",
                )

                pendentes = models.get_pagamentos_recorrentes_pendentes(
                    date(2026, 10, 7)
                )
                meses = [
                    pagamento["ano_mes"]
                    for pagamento in pendentes
                    if pagamento["descricao"] == "Aluguel retroativo"
                ]
                self.assertEqual(
                    meses,
                    [
                        "2026-06", "2026-07", "2026-08",
                        "2026-09", "2026-10",
                    ],
                )

                conn = database.get_db_connection()
                self.assertEqual(
                    conn.execute(
                        "SELECT COUNT(*) FROM lancamentos "
                        "WHERE descricao = 'Aluguel retroativo';"
                    ).fetchone()[0],
                    0,
                )
                conn.close()

                response = self.app.get("/recorrentes?ano=2026")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b'name="mes_inicio"', response.data)
                self.assertIn(
                    f'value="{date.today():%Y-%m}"'.encode(), response.data
                )
                self.assertIn(b'aria-label="Aluguel retroativo: 06/2026"', response.data)
                self.assertIn(b'aria-label="Aluguel retroativo: 10/2026"', response.data)

    def test_30_novo_recorrente_aceita_mes_de_inicio_no_passado(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                response = self.app.post(
                    "/recorrentes/novo",
                    data={
                        "descricao": "Plano antigo",
                        "valor_mensal": "49,90",
                        "categoria": "Tecnologia",
                        "dia_cobranca": "10",
                        "metodo_pagamento": "Pix",
                        "mes_inicio": "2026-06",
                    },
                )
                self.assertEqual(response.status_code, 302)
                conn = database.get_db_connection()
                mes_inicio = conn.execute(
                    "SELECT mes_inicio FROM gastos_recorrentes "
                    "WHERE descricao = 'Plano antigo';"
                ).fetchone()["mes_inicio"]
                conn.close()
                self.assertEqual(mes_inicio, "2026-06")

    def test_31_filtro_mensal_de_lancamentos_respeita_formato_de_data(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                conn.execute(
                    "UPDATE configuracoes SET formato_data = 'extenso' WHERE id = 1;"
                )
                conn.commit()
                conn.close()

                response = self.app.get("/lancamentos")
                self.assertEqual(response.status_code, 200)
                self.assertIn(
                    f"janeiro de {date.today().year}".encode(), response.data
                )
                self.assertIn(b"Todos", response.data)
                self.assertNotIn(b'type="month"', response.data)

    def test_32_dashboard_mostra_utilizacao_dos_limites_do_mes(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                conn = database.get_db_connection()
                conn.execute("UPDATE limites_categoria SET limite_mensal = 0;")
                conn.commit()
                conn.close()
                models.update_limite_categoria("Alimentação", 500)
                models.update_limite_categoria("Saúde", 319.19)
                models.update_limite_categoria("Transporte", 300)
                models.add_lancamento(
                    date.today().isoformat(), "Consulta", "Saúde", 319.99, "", "Pix"
                )
                models.add_gasto_recorrente(
                    "Estacionamento", 300, "Transporte", date.today().day, "Pix"
                )

                resumo = models.get_resumo_mensal(date.today().strftime("%Y-%m"))
                self.assertEqual(resumo["limite_total_estipulado"], 1119.19)
                self.assertEqual(resumo["total_gastos"], 619.99)
                self.assertEqual(resumo["percentual_limite"], 55.4)

                response = self.app.get("/")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b'aria-label="Limite mensal utilizado"', response.data)
                self.assertIn(b'aria-valuenow="55.4"', response.data)

    def test_33_limite_do_cartao_inclui_recorrencia_e_barra_contrastante(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_cartao("Cartão recorrente", 1000, 25, 5, "#00aeef")
                conn = database.get_db_connection()
                cartao_id = conn.execute(
                    "SELECT id FROM cartoes WHERE nome = ?;",
                    ("Cartão recorrente",),
                ).fetchone()["id"]
                conn.close()
                models.add_gasto_recorrente(
                    "Assinatura do cartão", 125, "Tecnologia",
                    date.today().day, "Cartão de Crédito", cartao_id,
                )

                cartao = models.get_cartoes()[0]
                self.assertEqual(cartao["limite_comprometido"], 125)
                self.assertEqual(cartao["percentual_uso"], 12.5)

                response = self.app.get("/cartoes")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b'data-contrast-progress', response.data)
                self.assertIn(
                    b'aria-label="Limite utilizado do cart\xc3\xa3o Cart\xc3\xa3o recorrente"',
                    response.data,
                )

    def test_34_lancamentos_exibem_compra_mais_recente_primeiro(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                for dia, descricao in (
                    ("2026-10-01", "Compra outubro"),
                    ("2026-08-15", "Compra agosto"),
                    ("2026-09-20", "Compra setembro"),
                ):
                    models.add_lancamento(
                        dia, descricao, "Alimentação", 10, "", "Pix"
                    )

                response = self.app.get("/lancamentos")
                self.assertEqual(response.status_code, 200)
                pos_outubro = response.data.index(b"Compra outubro")
                pos_setembro = response.data.index(b"Compra setembro")
                pos_agosto = response.data.index(b"Compra agosto")
                self.assertLess(pos_outubro, pos_setembro)
                self.assertLess(pos_setembro, pos_agosto)
                self.assertIn(b'aria-sort="descending"', response.data)
                self.assertIn(b"01/10/2026", response.data)
                self.assertIn(b'placeholder="dd/mm/aaaa"', response.data)

                dashboard = self.app.get("/?mes=2026-10")
                self.assertEqual(dashboard.status_code, 200)
                self.assertIn(b"01/10/2026", dashboard.data)

    def test_39_limite_do_cartao_considera_faturas_a_vencer(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_cartao("Cartão teste", 1000, 25, 5)
                conn = database.get_db_connection()
                cartao_id = conn.execute(
                    "SELECT id FROM cartoes WHERE nome = ?;",
                    ("Cartão teste",),
                ).fetchone()["id"]
                conn.close()

                primeiro_dia_mes = date.today().replace(day=1)
                mes_anterior = database.add_months(primeiro_dia_mes, -1)
                dois_meses_atras = database.add_months(primeiro_dia_mes, -2)
                models.add_lancamento(
                    mes_anterior.isoformat(), "Fatura ainda não vencida",
                    "Alimentação", 200, "", "Cartão de Crédito", cartao_id,
                )
                models.add_lancamento(
                    dois_meses_atras.isoformat(), "Fatura já vencida",
                    "Alimentação", 50, "", "Cartão de Crédito", cartao_id,
                )

                cartao = models.get_cartoes()[0]
                self.assertEqual(cartao["limite_comprometido"], 200)
                self.assertEqual(cartao["limite_disponivel"], 800)

    def test_40_preset_de_cor_xp_disponivel_no_cadastro_e_edicao(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                response = self.app.get("/cartoes")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.data.count(b'option value="#000000">XP'), 2)

    def test_38_datas_de_lancamentos_sao_exibidas_em_dia_mes_ano(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_lancamento(
                    "2026-05-10", "Compra de maio", "Alimentação", 10, "", "Pix"
                )

                lancamentos = self.app.get("/lancamentos")
                self.assertEqual(lancamentos.status_code, 200)
                self.assertIn(b"10/05/2026", lancamentos.data)
                self.assertIn(b'placeholder="dd/mm/aaaa"', lancamentos.data)
                self.assertIn(b'id="edit_data_display"', lancamentos.data)

                dashboard = self.app.get("/?mes=2026-05")
                self.assertEqual(dashboard.status_code, 200)
                self.assertIn(b"10/05/2026", dashboard.data)

    def test_35_modal_edicao_tem_seletor_de_cartao_oculto_ate_selecionar_credito(self):
        response = self.app.get("/lancamentos")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'id="edit_metodo_pagamento"', response.data)
        self.assertIn(
            b'<div class="col-12 d-none" id="edit_cartao_container">',
            response.data,
        )
        self.assertIn(
            b'<option value="">Selecione o cart\xc3\xa3o de cr\xc3\xa9dito</option>',
            response.data,
        )

    def test_36_editar_compra_parcelada_atualiza_cadastro_e_cronograma(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_compra_parcelada(
                    "Sofá", "2026-01-10", 600, 60, 3,
                    "Moradia", "Boleto / Carnê", None, "Compra original",
                )
                conn = database.get_db_connection()
                compra_id = conn.execute(
                    "SELECT id FROM compras_parceladas WHERE descricao = 'Sofá';"
                ).fetchone()["id"]
                conn.close()
                response = self.app.get("/parcelas")
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"data-edit-compra-parcelada", response.data)
                self.assertIn(
                    f'data-edit-id="{compra_id}"'.encode(), response.data
                )

                response = self.app.post(
                    f"/parcelas/editar/{compra_id}",
                    data={
                        "descricao": "Sofá novo",
                        "data_compra": "2026-02-15",
                        "valor_total": "720,00",
                        "entrada": "120,00",
                        "num_parcelas": "4",
                        "categoria": "Outros",
                        "metodo_pagamento": "Boleto / Carnê",
                        "observacao": "Atualizado",
                    },
                )
                self.assertEqual(response.status_code, 302)

                conn = database.get_db_connection()
                compra = conn.execute(
                    "SELECT descricao, data_compra, valor_total, entrada, "
                    "num_parcelas, valor_parcela, categoria, observacao "
                    "FROM compras_parceladas WHERE id = ?;",
                    (compra_id,),
                ).fetchone()
                parcelas = conn.execute(
                    "SELECT numero_parcela, total_parcelas, ano_mes, valor, "
                    "data_vencimento, status FROM parcelas_detalhe "
                    "WHERE compra_id = ? ORDER BY numero_parcela;",
                    (compra_id,),
                ).fetchall()
                entrada = conn.execute(
                    "SELECT descricao, categoria, valor, data FROM lancamentos "
                    "WHERE observacao = ?;",
                    (f"Entrada da compra parcelada #{compra_id}",),
                ).fetchone()
                conn.close()

                self.assertEqual(compra["descricao"], "Sofá novo")
                self.assertEqual(compra["data_compra"], "2026-02-15")
                self.assertEqual(compra["valor_total"], 720)
                self.assertEqual(compra["entrada"], 120)
                self.assertEqual(compra["num_parcelas"], 4)
                self.assertEqual(compra["valor_parcela"], 150)
                self.assertEqual(compra["categoria"], "Outros")
                self.assertEqual(compra["observacao"], "Atualizado")
                self.assertEqual(len(parcelas), 4)
                self.assertEqual(parcelas[0]["ano_mes"], "2026-02")
                self.assertEqual(parcelas[0]["data_vencimento"], "2026-02-15")
                self.assertEqual(parcelas[-1]["valor"], 150)
                self.assertTrue(all(p["status"] == "Pendente" for p in parcelas))
                self.assertEqual(entrada["descricao"], "Entrada: Sofá novo")
                self.assertEqual(entrada["categoria"], "Outros")
                self.assertEqual(entrada["valor"], 120)
                self.assertEqual(entrada["data"], "2026-02-15")

    def test_37_edicao_de_compra_com_parcela_paga_preserva_pagamento(self):
        with tempfile.TemporaryDirectory() as pasta:
            with patch.object(
                database, "DB_PATH", os.path.join(pasta, "finance.db")
            ):
                database.init_db()
                models.add_compra_parcelada(
                    "Notebook", "2026-01-10", 600, 0, 3,
                    "Tecnologia", "Boleto / Carnê", None,
                )
                conn = database.get_db_connection()
                compra_id = conn.execute(
                    "SELECT id FROM compras_parceladas WHERE descricao = 'Notebook';"
                ).fetchone()["id"]
                parcela_id = conn.execute(
                    "SELECT id FROM parcelas_detalhe WHERE compra_id = ? "
                    "ORDER BY numero_parcela LIMIT 1;",
                    (compra_id,),
                ).fetchone()["id"]
                conn.close()
                self.app.post(
                    f"/parcelas/{parcela_id}/pagar",
                    data={"metodo_pagamento": "Pix"},
                )

                response = self.app.post(
                    f"/parcelas/editar/{compra_id}",
                    data={
                        "descricao": "Notebook atualizado",
                        "data_compra": "2026-01-10",
                        "valor_total": "600",
                        "entrada": "0",
                        "num_parcelas": "4",
                        "categoria": "Tecnologia",
                        "metodo_pagamento": "Boleto / Carnê",
                    },
                    follow_redirects=True,
                )
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"Desfa\xc3\xa7a os pagamentos", response.data)

                conn = database.get_db_connection()
                compra = conn.execute(
                    "SELECT descricao, num_parcelas FROM compras_parceladas "
                    "WHERE id = ?;",
                    (compra_id,),
                ).fetchone()
                pagamento = conn.execute(
                    "SELECT p.status, p.lancamento_id, l.descricao "
                    "FROM parcelas_detalhe p LEFT JOIN lancamentos l "
                    "ON l.id = p.lancamento_id WHERE p.id = ?;",
                    (parcela_id,),
                ).fetchone()
                conn.close()
                self.assertEqual(compra["descricao"], "Notebook")
                self.assertEqual(compra["num_parcelas"], 3)
                self.assertEqual(pagamento["status"], "Pago")
                self.assertIsNotNone(pagamento["lancamento_id"])
                self.assertEqual(pagamento["descricao"], "Notebook")

                response = self.app.post(
                    f"/parcelas/editar/{compra_id}",
                    data={
                        "descricao": "Notebook atualizado",
                        "data_compra": "2026-01-10",
                        "valor_total": "600",
                        "entrada": "0",
                        "num_parcelas": "3",
                        "categoria": "Outros",
                        "metodo_pagamento": "Boleto / Carnê",
                        "observacao": "Categoria e nome revisados",
                    },
                )
                self.assertEqual(response.status_code, 302)
                conn = database.get_db_connection()
                pagamento = conn.execute(
                    "SELECT l.descricao, l.categoria FROM parcelas_detalhe p "
                    "JOIN lancamentos l ON l.id = p.lancamento_id "
                    "WHERE p.id = ?;",
                    (parcela_id,),
                ).fetchone()
                conn.close()
                self.assertEqual(pagamento["descricao"], "Notebook atualizado")
                self.assertEqual(pagamento["categoria"], "Outros")


if __name__ == "__main__":
    unittest.main()
